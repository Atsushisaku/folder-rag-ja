from pathlib import Path

import docx

from rag_chat.bm25 import BM25, tokenize
from rag_chat.chunking import split_text
from rag_chat.loaders import load_file
from rag_chat.index import scan_files
from rag_chat.ollama_client import has_model


def test_tokenize_mixes_bigrams_and_words():
    toks = tokenize("VPNの接続方法")
    assert "vpn" in toks
    assert "接続" in toks and "方法" in toks


def test_tokenize_normalizes_fullwidth():
    assert tokenize("ＶＰＮ") == ["vpn"]


def test_bm25_ranks_matching_doc_first():
    bm = BM25(["経費の精算は20日締め", "会議室はカレンダーで予約", "VPNの初回設定"])
    scores = bm.scores("会議室の予約")
    assert scores.index(max(scores)) == 1


def test_split_text_respects_size_and_keeps_content():
    text = "\n\n".join(f"段落{i}" + "あ" * 150 for i in range(10))
    chunks = split_text(text, size=400, overlap=50)
    assert len(chunks) > 1
    assert all(len(c) <= 400 for c in chunks)
    assert "段落9" in chunks[-1]


def test_split_text_long_paragraph_overlaps():
    chunks = split_text("い" * 1000, size=300, overlap=50)
    assert all(len(c) <= 300 for c in chunks)
    assert sum(len(c) for c in chunks) >= 1000


def test_markdown_sections_have_heading_path(tmp_path: Path):
    p = tmp_path / "m.md"
    p.write_text("# 大\n前文\n## 中\n本文A\n```\n# コード内は見出しではない\n```\n## 次\n本文B\n", encoding="utf-8")
    secs = load_file(p)
    assert [s.location for s in secs] == ["大", "大 > 中", "大 > 次"]
    assert "コード内" in secs[1].text


def test_docx_sections(tmp_path: Path):
    d = docx.Document()
    d.add_heading("手順", level=1)
    d.add_paragraph("ログインする")
    p = tmp_path / "w.docx"
    d.save(p)
    secs = load_file(p)
    assert secs[0].location == "手順"
    assert "ログイン" in secs[0].text


def test_scan_files_skips_hidden_and_temp(tmp_path: Path):
    (tmp_path / "a.md").write_text("x", encoding="utf-8")
    (tmp_path / "~$b.docx").write_text("x", encoding="utf-8")
    (tmp_path / ".hidden").mkdir()
    (tmp_path / ".hidden" / "c.md").write_text("x", encoding="utf-8")
    (tmp_path / "d.url").write_text("x", encoding="utf-8")
    assert [p.name for p in scan_files(tmp_path)] == ["a.md"]


def test_has_model_latest_tag():
    assert has_model(["bge-m3:latest"], "bge-m3")
    assert has_model(["qwen3:4b"], "qwen3:4b")
    assert not has_model(["qwen3:8b"], "qwen3:4b")


def test_tokenize_drops_hiragana_bigrams():
    toks = tokenize("風邪で休むときはどうすればいい")
    assert "風邪" in toks
    assert not any(t in toks for t in ("とき", "どう", "邪で", "いい"))


def test_tokenize_keeps_kanji_between_hiragana():
    assert tokenize("立て替えたお金") == ["立", "替", "金"]


def test_cited_sources_only_cited_and_in_range():
    from rag_chat.answer import cited_sources
    hits = [{"path": f"f{i}.md"} for i in range(1, 6)]
    got = cited_sources("手順です [2]。補足 ［4］ と [2]、範囲外 [9]", hits)
    assert [(n, h["path"]) for n, h in got] == [(2, "f2.md"), (4, "f4.md")]
    assert cited_sources("資料には見当たりません", hits) == []
