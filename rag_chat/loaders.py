"""ファイルからテキストを取り出し、出典の位置 (ページ・見出し) 付きのセクションにする。"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTS = {".md", ".markdown", ".txt", ".pdf", ".docx"}

_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


@dataclass
class Section:
    text: str
    location: str  # 例: "p.3" / "手順 > 申請方法"


def load_file(path: Path) -> list[Section]:
    ext = path.suffix.lower()
    if ext in (".md", ".markdown"):
        return _load_markdown(path.read_text(encoding="utf-8", errors="replace"))
    if ext == ".txt":
        return [Section(path.read_text(encoding="utf-8", errors="replace"), "")]
    if ext == ".pdf":
        return _load_pdf(path)
    if ext == ".docx":
        return _load_docx(path)
    raise ValueError(f"未対応の形式: {path.suffix}")


def _load_markdown(text: str) -> list[Section]:
    """見出しごとに区切る。location は見出しの階層を " > " でつないだもの。"""
    sections: list[Section] = []
    stack: list[tuple[int, str]] = []
    buf: list[str] = []
    in_code = False

    def flush() -> None:
        body = "\n".join(buf).strip()
        if body:
            sections.append(Section(body, " > ".join(t for _, t in stack)))
        buf.clear()

    for line in text.splitlines():
        if line.lstrip().startswith("```"):
            in_code = not in_code
        m = None if in_code else _MD_HEADING.match(line)
        if m:
            flush()
            level = len(m.group(1))
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, m.group(2)))
        else:
            buf.append(line)
    flush()
    return sections


def _load_pdf(path: Path) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    out = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            out.append(Section(text, f"p.{i}"))
    return out


def _load_docx(path: Path) -> list[Section]:
    """Word の見出しスタイル (Heading n / 見出し n) で区切る。"""
    import docx

    document = docx.Document(str(path))
    sections: list[Section] = []
    stack: list[tuple[int, str]] = []
    buf: list[str] = []

    def flush() -> None:
        body = "\n".join(buf).strip()
        if body:
            sections.append(Section(body, " > ".join(t for _, t in stack)))
        buf.clear()

    for para in document.paragraphs:
        style = (para.style.name if para.style is not None else "") or ""
        m = re.match(r"(?:Heading|見出し)\s*(\d)", style)
        text = para.text.strip()
        if m and text:
            flush()
            level = int(m.group(1))
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, text))
        elif text:
            buf.append(text)
    for table in document.tables:
        for row in table.rows:
            buf.append(" | ".join(c.text.strip() for c in row.cells))
    flush()
    return sections
